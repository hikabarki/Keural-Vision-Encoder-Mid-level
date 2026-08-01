"""
Rewrite Keural_Mid_Technical_Documentation_Report.docx from the audited original.

Every change below is traceable to a finding from the verification pass that
checked all 147 factual claims in the document against the source code and the
result files. Claim ids (C5, C23, ...) are quoted in the comments.

The document is COPIED and edited in place so the letterhead, styles, equation
images and page geometry are inherited untouched. Only text that the audit found
wrong or outdated is replaced, plus one new results section and an expanded
reference list.
"""
from __future__ import annotations

import copy
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "docs/reports/Keural_Mid_Technical_Documentation_Report.docx"
DST = ROOT / "docs/reports/Keural_Mid_Technical_Documentation_Report_v2.docx"
FIGS = ROOT / "docs/figures"
WIDTH = Inches(6.0)

shutil.copy(SRC, DST)
doc = Document(DST)
P = doc.paragraphs


def settext(idx: int, text: str) -> None:
    """Replace a paragraph's text, keeping its style and first-run formatting."""
    p = P[idx]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def new_par(after_idx: int, text: str, style: str = "Normal"):
    """Insert a new paragraph immediately AFTER the given index."""
    src = P[after_idx]
    q = doc.add_paragraph(text, style=style)
    src._p.addnext(q._p)
    return q


# =============================================================================
# FRONT MATTER  (C2)
# =============================================================================
settext(14, "Status: Stage 1 (Contrastive Pre-training) — Complete; "
            "Stage 2 (Vision-Language Alignment) — Complete "
            "(2a, 2b English, 2b Bilingual); Stage 2c (OCR) — In Progress")
settext(15, "Document date: 2026-08-01")

# =============================================================================
# 1. EXECUTIVE SUMMARY  (C5 critical, C12, C13)
# =============================================================================
# C5: the token COUNT is not content-dependent. _resolve_default_budget scales it
# from pixel count alone, and both Stage 1 and Stage 2 held it fixed at 512. The
# saliency map decides WHERE tokens are placed, not HOW MANY there are.
settext(19,
    "Keural Mid is a 183.6-million-parameter vision encoder (183,575,297 parameters) "
    "trained entirely from random initialization — it is not a fine-tune of CLIP, ViT, "
    "or DINO. It combines three mechanisms that, to our knowledge, do not appear "
    "together in an existing open encoder. First, an Adaptive Token Budget (ATB) "
    "tokenizer that places a content-adaptive distribution of tokens: positions are "
    "chosen by a learned saliency map through differentiable Gumbel-Top-K sampling "
    "with non-maximum suppression, under a token budget that is set at call time or "
    "auto-scaled with input resolution (default 512, maximum 2,048). The budget "
    "governs how many tokens are emitted; the saliency map governs where they are "
    "placed. Second, Hierarchical Concept Tokenization (HCT), which tags every token "
    "with a semantic level — global, region, or detail — and gives each level its own "
    "learned embedding and constructor. Third, continuous 2D rotary position encoding "
    "computed from each token's real normalized (cx, cy) coordinates rather than a "
    "fixed patch-grid index, which is what allows the encoder to accept any input "
    "resolution that is a multiple of 16 without interpolating a positional table.")

settext(21,
    "This report documents the completion of Stage 1 (contrastive pre-training) and "
    "Stage 2 (vision-language alignment), including full architectural specification, "
    "training methodology, mathematical formulation of every loss term, and evaluation "
    "results for both stages. Stage 1 ran for 18,000 optimizer steps on 20M image-text "
    "pairs, reducing total training loss from 21.56 to 2.87 and reaching 37.2% "
    "zero-shot ImageNet Top-1, 42.9%/42.2% Flickr30K R@1, and 75.0% CIFAR-100 linear "
    "probe. Stage 2 attached the frozen encoder to Qwen2.5-7B-Instruct through a "
    "trained projector, producing a bilingual (English and Korean) vision-language "
    "model. The headline Stage 2 result is the measured contribution of the vision "
    "pipeline itself: against an identical model with the visual tokens removed, the "
    "sighted model gains +24.6 accuracy points on POPE (+46.6 F1) and +11.8 points on "
    "VQAv2. That gap cannot be produced by the language model alone and is the "
    "clearest available evidence that the encoder carries real visual information.")

# C12: CLIP ViT-B/16 (~86M) and SigLIP ViT-B/16 (~93M) are SMALLER than Keural Mid,
# not larger. The original wording said "larger ... peers", which is backwards.
settext(22,
    "Accuracy trails peers that are smaller in parameter count but far more "
    "extensively trained — CLIP ViT-B/16 (~86M parameters, 400M pairs) and SigLIP "
    "ViT-B/16 (~93M, WebLI) — and the gap is attributed to training budget rather "
    "than an architectural deficiency: Stage 1 saw 20M pairs over 18,000 steps and "
    "its loss was still declining when the run ended. This report also lays out the "
    "mathematical foundations for the two remaining alignment stages, Direct "
    "Preference Optimization (Stage 3) and Reinforcement Learning from Human "
    "Feedback (Stage 4), neither of which has yet been run.")

# =============================================================================
# 2. ARCHITECTURE  (C23 critical, C37)
# =============================================================================
# C23: there is no ROI-Align anywhere in the model. token_constructor.py pools with
# per-region Gaussian soft weights via einsum. roi_align_size is dead config.
settext(32,
    "Given the stem feature map and the saliency map, the ATB tokenizer selects "
    "spatially diverse token positions via a differentiable Gumbel-softmax top-K "
    "sampler with non-maximum suppression (NMS radius 3 for region-level, 1 for "
    "detail-level). A feature vector is then extracted per selected position by "
    "differentiable weighted pooling — not by ROI-Align. Each position receives a "
    "Gaussian soft-attention map over the stem grid (sigma = smoothing kernel / 4, "
    "i.e. 1.75 at region level and 0.75 at detail level, normalized to sum to one), "
    "and the token is the weighted average of stem features under that map, computed "
    "as a single einsum. Region and detail levels use separate constructor weights, "
    "each a residual MLP (LayerNorm, Linear 768→1536, GELU, Linear 1536→768, "
    "LayerNorm). Soft pooling is used in place of hard ROI-Align precisely because "
    "the sampled indices must remain differentiable so that gradients reach the "
    "saliency predictor. The default token budget is 512 tokens (maximum 2,048), "
    "split 5% global / 25% region / 70% detail — 26 / 128 / 358 tokens at the default "
    "budget — with floors of 8 global and 16 region tokens.")

# C37: the same ROI-align error is repeated in the Mermaid source, and the loss box
# omits that the saliency term is scaled by lambda_sal * beta_sal, not lambda_sal.
mermaid = P[43].text
mermaid = mermaid.replace("Gumbel top-K + NMS + ROI-align",
                          "Gumbel top-K + NMS + Gaussian soft-weight pooling")
mermaid = mermaid.replace("L_SigLIP + λ_hct·L_HCT + λ_sal·β·L_saliency + L_distill",
                          "L_SigLIP + λ_hct·L_HCT + λ_sal·β_sal·L_saliency + λ_distill·L_distill")
settext(43, mermaid)

# =============================================================================
# 3. DATA PIPELINE  (C38, C39, C42, C44, C45, C49)
# =============================================================================
# C42: both declared buckets are 1:1, and _assign_bucket tie-breaks on strict '<',
# so bucket 0 (384) always won. The 448 bucket was never populated.
# C44: Stage 1 DOES augment - RandomResizedCrop + flip + ColorJitter.
settext(50,
    "Two square resolution buckets (384×384 and 448×448) are declared in the training "
    "configuration, but in practice all Stage 1 training ran at 384×384: bucket "
    "assignment selects on aspect ratio alone, and because both buckets are 1:1 the "
    "strict-less-than tie-break always resolves to the first bucket. The 448×448 "
    "bucket was never populated. Preprocessing applies RandomResizedCrop to the "
    "bucket size (scale 0.75–1.0, ratio 3/4–4/3, bicubic interpolation), "
    "RandomHorizontalFlip (p = 0.5), and ColorJitter (brightness 0.1, contrast 0.1, "
    "saturation 0.1, hue 0.03), followed by normalization with ImageNet statistics "
    "(mean [0.485, 0.456, 0.406], std [0.229, 0.224, 0.225]). Augmentation is "
    "therefore deliberately light rather than absent; no heavy policies (RandAugment, "
    "mixup, CutMix, RandomErasing) are used, and those remain a candidate for a "
    "longer future run.")

settext(52,
    "A BucketedSampler groups sample indices by resolution bucket so that every batch "
    "has a uniform spatial shape, additionally sorts within a bucket by source shard "
    "for I/O locality with window-local shuffling, and shards the resulting batch list "
    "per-rank for distributed data-parallel training. It does not pad; all images in a "
    "batch are already the same size by construction. Effective batch size is 3,072 "
    "(768 samples per GPU × 2 gradient-accumulation steps × 2 GPUs), with distributed "
    "all_gather used to assemble in-batch negatives across ranks for the contrastive "
    "loss.")

# C49: the original asserted a causal claim ("did not affect throughput ... not the
# bottleneck") that was never profiled. State the measurements, not the attribution.
settext(53,
    "Operational note: dataloader worker count was reduced from 12 to 6 per rank "
    "during this training run after a host-memory OOM (each worker held a ~9.7 GB "
    "copy of the sample index; see Section 4.3 for training continuity details). The "
    "resumed run sustained approximately 25.0 s/step (about 123 images/s at effective "
    "batch 3,072), compared with approximately 30.2 s/step during the healthy phase of "
    "the first run. No profiling was performed to attribute the difference, so no "
    "causal claim is made about the data pipeline being or not being the bottleneck.")

# =============================================================================
# 4. EQUATION CAPTIONS  (C59, C65)
# =============================================================================
# C59: the learnable scale enters as exp(t'), not as tau directly.
settext(60,
    "Eq. 1 — scaled, biased cosine similarity between image embedding x_i and text "
    "embedding y_j. The scale enters as exp(t′) where t′ (logit_scale) and the bias b "
    "(logit_bias) are both learnable scalars, initialized t′ = 2.30 and b = −10.0, "
    "and reaching t′ ≈ 4.34 (effective scale ≈ 77) and b ≈ −9.7 by step 18,000.")

# C65: the saliency term carries lambda_sal * beta_sal = 0.5 * 0.005 = 0.0025.
settext(78,
    "Eq. 7 — full Stage 1 loss, combining all four terms above. The configured "
    "weights are λ_hct = 1.0, λ_sal = 0.5 with β_sal = 0.005 (an effective saliency "
    "weight of 0.0025), and λ_distill = 0.5, split evenly at 0.25 per teacher. Note "
    "that the saliency regularizer is scaled by the product λ_sal · β_sal, not by "
    "λ_sal alone.")

# =============================================================================
# 4.3 SALIENCY STABILITY  (C93 critical)
# =============================================================================
# C93: the original "~50% by step 13,000-14,000" contradicts the report's own
# Figure 6, which shows ~49-52% at steps 10,000-11,000 and ~65% at step 14,000.
settext(101,
    "The learned saliency map produces sharp, stable contours that track each image's "
    "dominant subject from mid-training through the final checkpoint. Quantitatively, "
    "the continuous saliency map stabilizes quickly — cosine similarity to the final "
    "checkpoint reaches roughly 0.83–0.90 within the first few thousand steps — but "
    "the discrete sampled token positions are considerably less stable. Jaccard "
    "overlap with the final checkpoint's token positions is only about 49–52% at "
    "steps 10,000–11,000, is still about 65% at step 14,000, and reaches 73% by step "
    "16,000. Exact token placement therefore continued to shift meaningfully late in "
    "training even where the saliency field appeared visually converged (values read "
    "from Figure 6, n = 8 images).")

doc.save(DST)
print(f"pass 1 complete -> {DST}")
print("  front matter, executive summary, architecture, data pipeline,")
print("  equation captions and saliency stability corrected")
